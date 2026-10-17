from docx import Document
from pathlib import Path
import json,re,csv
B=Path(__file__).resolve().parent; d=Document(B/'fuentes'/'PIC_Innovacion_Centro_26-27.docx')
objs=[
('1','Impulsar la creación de una cultura innovadora en el centro'),
('2','Incorporar nuevas metodologías de enseñanza para favorecer la innovación pedagógica en Formación Profesional'),
('3','Fomentar la realización de proyectos de innovación'),
('4','Explorar e integrar la innovación tecnológica y la digitalización en los procesos de enseñanza-aprendizaje y en la organización del centro'),
('5','Promover la colaboración con las empresas y el entorno productivo'),
('6','Impulsar la transferencia de conocimiento y la difusión de las buenas prácticas de innovación')]
cat={'id':'PIC_2026_2027_INNOVACION_C2','version':'1.0.0','estado':'transcripcion_para_revision','fecha_corte':'2026-10-17','fuente':'PIC_Innovacion_Centro_26-27.docx (modelo facilitado para cumplimentacion en Laycos)','reglas':{'objetivos':6,'acciones_minimas_por_objetivo':1,'acciones_obligatorias':['1.1','1.2'],'otras_acciones_maximo':3,'proyectos_maximo_formulario':6,'dafo_por_cuadrante_minimo':1,'dafo_por_cuadrante_maximo':3,'reuniones_CICLOPE_minimo':4,'plazo_pic':'2026-10-30','plazo_memoria':'2027-06-18','riesgo_numerico_positivo':True,'logrado_actualizable_solo_en_seguimiento':True},'objetivos':[]}
for n,(idobj,title) in enumerate(objs):
 tbl=d.tables[10+n]; acts={}
 for row in tbl.rows[1:]:
  cols=[c.text.strip() for c in row.cells]
  action=cols[1]; indicator=cols[2]
  code=re.match(r'^(\d+\.\d+)\.',action)
  if not code: raise ValueError((n,action))
  code=code.group(1)
  if code not in acts:
   acts[code]={'codigo':code,'titulo':re.sub(r'^\d+\.\d+\.\s*','',action).replace('\n[OBLIGATORIA]','').strip(),'obligatoria':('[OBLIGATORIA]' in action),'indicadores':[]}
  v='booleano' if '(sí / no)' in indicator else 'porcentaje' if '(porcentaje)' in indicator else 'numero' if '(número)' in indicator else 'escala_1_5' if '(escala' in indicator else 'revisar_unidad'
  m=re.search(r'mínimo\s+(\d+)',indicator)
  acts[code]['indicadores'].append({'id':f'PIC-{code}-I{len(acts[code]["indicadores"])+1}','descripcion':indicator,'tipo':v,'minimo_normativo':int(m.group(1)) if m else None,'metodo_calculo':'CONFIGURAR_Y_VALIDAR_CON_C2','fuente_evidencia':'pendiente_definicion','conteo':'UNICO_SALVO_REGLA_ESPECIFICA'})
 cat['objetivos'].append({'codigo':idobj,'titulo':title,'acciones':list(acts.values())})
cat['familias_profesionales']=[r.cells[0].text.strip().lstrip('· ').strip() for r in d.tables[29].rows if r.cells[0].text.strip()]
cat['ods_2030']=[r.cells[0].text.strip().lstrip('· ').strip() for r in d.tables[30].rows if r.cells[0].text.strip()]
cat['ods_canarios']=[r.cells[0].text.strip().lstrip('· ').strip() for r in d.tables[31].rows if r.cells[0].text.strip()]
cat['cargos_ciclope']=[r.cells[0].text.strip().lstrip('· ').strip() for r in d.tables[32].rows if r.cells[0].text.strip()]
cat['dafo_categorias']={}
for name,tbl in zip(['debilidades','fortalezas','amenazas','oportunidades'],d.tables[33:37]):
 cat['dafo_categorias'][name]=[{'categoria':row.cells[0].text.strip(),'opciones': [x.strip(' ·\n') for x in row.cells[1].text.strip().split('\n') if x.strip(' ·\n')]} for row in tbl.rows[1:] if row.cells[0].text.strip()]
with (B/'config/pic_2026_2027_catalogo.json').open('w',encoding='utf8') as f:json.dump(cat,f,ensure_ascii=False,indent=2)
with (B/'config/pic_2026_2027_acciones.csv').open('w',encoding='utf8',newline='') as f:
 w=csv.writer(f);w.writerow(['objetivo_codigo','objetivo','accion_codigo','accion','obligatoria','indicador_id','indicador','tipo','minimo_normativo'])
 for obj in cat['objetivos']:
  for a in obj['acciones']:
   for ind in a['indicadores']:
    w.writerow([obj['codigo'],obj['titulo'],a['codigo'],a['titulo'],int(a['obligatoria']),ind['id'],ind['descripcion'],ind['tipo'],ind['minimo_normativo'] if ind['minimo_normativo'] is not None else ''])
print('Objetivos',len(cat['objetivos']),'acciones',sum(len(o['acciones']) for o in cat['objetivos']),'indicadores',sum(len(a['indicadores']) for o in cat['objetivos'] for a in o['acciones']))
for k,v in cat['dafo_categorias'].items():print(k,'categorias',len(v),'opciones',[len(e['opciones']) for e in v])
